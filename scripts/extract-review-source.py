"""Read the supplied Word bank without editing it; emit draft-only JSON."""
import json
import re
import sys
from docx import Document

doc = Document(sys.argv[1])
questions = []
day = None
for p in doc.paragraphs:
    text = p.text.strip()
    if re.match(r'Day [12] Passport Trivia', text):
        day = int(text[4])
    match = re.match(r'Question (\d+)\s+(.+?)\s+(Easy|Medium|Hard)$', text)
    if match:
        questions.append(dict(id=f'd{day}q{match[1]}', day=day, order=int(match[1]), category=match[2], difficulty=match[3], question='', options=[]))
    elif questions and re.match(r'^[A-D]\. ', text):
        questions[-1]['options'].append(text[3:])
    elif questions and not questions[-1]['question']:
        questions[-1]['question'] = text
tables = [t for t in doc.tables if t.cell(0, 0).text == 'Answer and explanation']
assert len(tables) == len(questions) == 24
for q, t in zip(questions, tables):
    answer = t.cell(0, 1).text.split('\n', 1)
    q.update(correctOption=ord(answer[0][0])-65, explanation=answer[1], originalVisual=t.cell(1, 1).text)
    assert len(q['options']) == 4
print(json.dumps(questions, ensure_ascii=False, indent=2))
